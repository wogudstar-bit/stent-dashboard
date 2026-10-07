import streamlit as st
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.gridspec as gridspec
import scipy.stats as stats
from statsmodels.stats.diagnostic import normal_ad
import io
import re

st.set_page_config(page_title="스텐트 성능 분석", layout="wide")
st.title("📊 스텐트 자동 분석 & 데이터 생성 대시보드")

tab1, tab2 = st.tabs(["📊 데이터 분석기", "🎲 임의 데이터 생성기"])

# =========================================================
# 탭 1: 기존 분석기 코드 (대분류 시험명 기준 그룹핑 적용)
# =========================================================
with tab1:
    pink_fill = PatternFill(start_color="FFC0CB", end_color="FFC0CB", fill_type="solid")

    def parse_criteria(crit_str):
        if crit_str is None: return "<=", 0.0
        crit_str = str(crit_str).replace(',', '').strip()
        if '~' in crit_str:
            parts = crit_str.split('~')
            if len(parts) == 2:
                try:
                    return "~", (float(parts[0].strip()), float(parts[1].strip()))
                except ValueError:
                    pass
        match = re.match(r"([<>]=?)\s*([\d\.]+)", crit_str)
        if match:
            return match.group(1), float(match.group(2))
        try:
            return "<=", float(crit_str)
        except ValueError:
            return "<=", 0.0

    def check_data_pass(value, operator, limit):
        if operator == '~': return limit[0] <= value <= limit[1]
        if operator == '<=': return value <= limit
        if operator == '<': return value < limit
        if operator == '>=': return value >= limit
        if operator == '>': return value > limit
        return True

    def create_minitab_plot(data, model_name):
        n = len(data)
        mean = np.mean(data)
        std = np.std(data, ddof=1)
        var = np.var(data, ddof=1)
        skew = stats.skew(data, bias=False)
        kurt = stats.kurtosis(data, bias=False)
        stat_ad, p_val_ad = normal_ad(data)
        minimum = np.min(data)
        q1 = np.percentile(data, 25)
        median = np.median(data)
        q3 = np.percentile(data, 75)
        maximum = np.max(data)
        ci_mean = stats.t.interval(0.95, df=n-1, loc=mean, scale=std/np.sqrt(n))
        ci_median = (median - 1.57*(q3-q1)/np.sqrt(n), median + 1.57*(q3-q1)/np.sqrt(n))
        ci_std = (std * np.sqrt((n-1)/stats.chi2.ppf(0.975, n-1)), std * np.sqrt((n-1)/stats.chi2.ppf(0.025, n-1)))

        fig = plt.figure(figsize=(10, 6))
        gs = gridspec.GridSpec(3, 2, width_ratios=[2.5, 1], height_ratios=[3, 1, 1.2])
        fig.suptitle(f"Summary Report for {model_name}", fontsize=14)

        ax1 = fig.add_subplot(gs[0, 0])
        ax1.hist(data, bins=5, density=True, alpha=0.6, color='#5c8ccc', edgecolor='black')
        xmin, xmax = ax1.get_xlim()
        x = np.linspace(xmin, xmax, 100)
        p = stats.norm.pdf(x, mean, std)
        ax1.plot(x, p, 'darkred', linewidth=1.5)
        ax1.grid(True, alpha=0.3)
        ax1.set_yticks([])

        ax2 = fig.add_subplot(gs[1, 0], sharex=ax1)
        ax2.boxplot(data, vert=False, patch_artist=True, boxprops=dict(facecolor='#5c8ccc', color='black'))
        ax2.set_yticks([])
        ax2.grid(True, alpha=0.3)

        ax3 = fig.add_subplot(gs[2, 0], sharex=ax1)
        ax3.errorbar((ci_mean[0]+ci_mean[1])/2, 1, xerr=(ci_mean[1]-ci_mean[0])/2, fmt='o', color='#0055a4', capsize=4)
        ax3.errorbar((ci_median[0]+ci_median[1])/2, 0.5, xerr=(ci_median[1]-ci_median[0])/2, fmt='o', color='#0055a4', capsize=4)
        ax3.set_yticks([0.5, 1])
        ax3.set_yticklabels(['Median', 'Mean'])
        ax3.set_ylim(0, 1.5)
        ax3.set_title("95% Confidence Intervals", fontsize=10)
        ax3.grid(True, alpha=0.3)

        ax4 = fig.add_subplot(gs[:, 1])
        ax4.axis('off')
        stats_text = (
            f"Anderson-Darling Normality Test\n"
            f"{'A-Squared':<15} {stat_ad:.2f}\n"
            f"{'P-Value':<15} {p_val_ad:.3f}\n\n"
            f"{'Mean':<15} {mean:.1f}\n"
            f"{'StDev':<15} {std:.1f}\n"
            f"{'N':<15} {n}\n\n"
            f"{'Minimum':<15} {minimum:.1f}\n"
            f"{'1st Quartile':<15} {q1:.1f}\n"
            f"{'Median':<15} {median:.1f}\n"
            f"{'3rd Quartile':<15} {q3:.1f}\n"
            f"{'Maximum':<15} {maximum:.1f}\n"
        )
        ax4.text(0.1, 0.95, stats_text, va='top', ha='left', fontsize=9, family='monospace')
        plt.tight_layout()
        
        img_buffer = io.BytesIO()
        plt.savefig(img_buffer, format='png', dpi=150)
        img_buffer.seek(0)
        return fig, img_buffer

    uploaded_file = st.file_uploader("검사 결과 엑셀 파일(빈 양식)을 이곳에 끌어다 놓으세요.", type=['xlsx'])
    if uploaded_file is not None:
        wb = openpyxl.load_workbook(uploaded_file)
        ws = wb.active
        
        try:
            k_val_one = float(ws['D2'].value)
            k_val_two = float(ws['E2'].value)
        except:
            st.warning("⚠ D2(단측) 또는 E2(양측) 셀에 K값이 비어있을 수 있습니다.")
            k_val_one, k_val_two = 0.0, 0.0
            
        blocks = []
        current_main_name = "미지정 시험" # 메인 시험명 추적용
        
        for r in range(1, ws.max_row + 15):
            c_val = ws.cell(row=r, column=1).value
            b_val = ws.cell(row=r, column=2).value
            
            if isinstance(c_val, str):
                c_str = c_val.strip()
                
                # 1. 메인 시험명(대분류) 자동 추적 로직
                # B열이 비어있고, 일반적인 통계 용어가 아닌 경우 대분류명으로 인식
                if b_val is None and c_str and c_str not in ["평균", "표준편차", "criteria", "K값", "p값", "BTR", "2A", "2R"]:
                    if re.match(r'^\d+[\.\)]', c_str): # "1. " 형태면 무조건 메인 제목 갱신
                        current_main_name = c_str
                    else: # 숫자가 없더라도 부제목(예: Body Diameter)이면 뒤에 이어붙임
                        if current_main_name != "미지정 시험" and not current_main_name.endswith(c_str):
                            current_main_name = f"{current_main_name} - {c_str}"
                        elif current_main_name == "미지정 시험":
                            current_main_name = c_str
                            
                # 2. 서브 데이터 묶음(BTR, 2A 등) 탐색 로직
                if c_str == "평균":
                    header_row = r - 1
                    while header_row > 1:
                        h_val = ws.cell(row=header_row, column=1).value
                        if h_val is not None:
                            h_str = str(h_val).strip()
                            if not h_str.replace('.', '', 1).isdigit():
                                break
                        header_row -= 1
                    
                    blocks.append({
                        'main_test_name': current_main_name,
                        'mean_row': r,
                        'header_row': header_row
                    })
                
        if not blocks:
            st.error("⚠️ A열에 '평균'이 포함되어 있는지 확인하세요.")
        else:
            # 추출된 블록들을 대분류(main_test_name) 기준으로 묶어주기 (Grouping)
            grouped_blocks = {}
            for block in blocks:
                m_name = block['main_test_name']
                if m_name not in grouped_blocks:
                    grouped_blocks[m_name] = []
                grouped_blocks[m_name].append(block)
                
            st.success(f"✅ 총 **{len(grouped_blocks)}개**의 시험 항목(대분류)과 **{len(blocks)}개**의 세부 묶음 스캔 완료!")
            
            # 대분류별로 묶어서 화면에 깔끔하게 출력
            for main_name, sub_blocks in grouped_blocks.items():
                st.markdown(f"## 📌 {main_name}") # 큰 제목 (예: 1. Deployment force)
                
                for block in sub_blocks:
                    mean_row = block['mean_row']
                    header_row = block['header_row']
                    test_name = str(ws.cell(row=header_row, column=1).value) # 서브 제목 (예: BTR)
                    
                    last_col = 2
                    while ws.cell(row=header_row, column=last_col).value is not None:
                        last_col += 1
                    last_col -= 1
                    
                    is_dimensional = True
                    crit_row = None
                    for r in range(mean_row, mean_row + 6):
                        c_val = ws.cell(row=r, column=1).value
                        if c_val is not None:
                            c_str = str(c_val).replace(" ", "").lower()
                            if "criteria" in c_str:
                                crit_row = r
                            if "k값" in c_str or "p값" in c_str:
                                is_dimensional = False
                                
                    if crit_row is None:
                        crit_row = mean_row + 4 

                    mode_text = "📏 치수 시험" if is_dimensional else "📈 성능 시험"
                    st.markdown(f"#### ↳ 🧪 {test_name} ({mode_text})")
                    
                    cols = st.columns(max(1, last_col - 1))
                    col_idx = 0
                    img_insert_col = max(7, last_col + 2)
                    
                    for col in range(2, last_col + 1):
                        model_name = ws.cell(row=header_row, column=col).value
                        if not model_name: continue
                            
                        raw_criteria = ws.cell(row=crit_row, column=col).value
                        
                        data_cells = []
                        data = []
                        for dr in range(header_row + 1, mean_row):
                            d_val = ws.cell(row=dr, column=col).value
                            if d_val is not None:
                                val_float = float(str(d_val).replace(',', ''))
                                data.append(val_float)
                                data_cells.append((dr, val_float))
                                
                        if not data: continue
                            
                        mean_val = np.mean(data)
                        std_val = np.std(data, ddof=1)
                        
                        if is_dimensional:
                            try:
                                base_val = float(str(raw_criteria).replace(',', ''))
                            except:
                                base_val = 0.0
                                
                            pct_val = ws.cell(row=crit_row+1, column=col).value
                            pct_float = 0.0
                            if isinstance(pct_val, str) and '%' in pct_val:
                                pct_float = float(pct_val.replace('%', '')) / 100.0
                            elif isinstance(pct_val, (int, float)):
                                if pct_val > 1: pct_float = pct_val / 100.0
                                else: pct_float = float(pct_val)
                                
                            lower_limit = base_val * (1 - pct_float)
                            upper_limit = base_val * (1 + pct_float)
                            
                            failed_count = 0
                            for dr, val_float in data_cells:
                                if not (lower_limit <= val_float <= upper_limit):
                                    ws.cell(row=dr, column=col).fill = pink_fill
                                    failed_count += 1
                                    
                            mean_cell = ws.cell(row=mean_row, column=col)
                            mean_cell.value = round(mean_val, 2)
                            is_mean_pass = (lower_limit <= mean_val <= upper_limit)
                            if not is_mean_pass:
                                mean_cell.fill = pink_fill
                                
                            ws.cell(row=mean_row+1, column=col).value = round(std_val, 2)
                            
                            with cols[col_idx % len(cols)]:
                                st.markdown(f"**{model_name}**")
                                if is_mean_pass:
                                    st.write(f"✅ **평균:** {round(mean_val, 2)}")
                                else:
                                    st.write(f"❌ **평균 미달:** {round(mean_val, 2)}")
                                    
                                limit_text = f"기준: {round(lower_limit, 2)} ~ {round(upper_limit, 2)}"
                                if failed_count > 0:
                                    st.error(f"❌ 불량 {failed_count}건 발생 ({limit_text})")
                                else:
                                    st.success(f"✅ 전수 통과 ({limit_text})")
                            
                        else:
                            operator, limit = parse_criteria(raw_criteria)
                            
                            for dr, val_float in data_cells:
                                if not check_data_pass(val_float, operator, limit):
                                    ws.cell(row=dr, column=col).fill = pink_fill
                                    
                            stat, p_val = normal_ad(np.array(data))
                            
                            if operator == '~':
                                ltl = mean_val - (k_val_two * std_val)
                                utl = mean_val + (k_val_two * std_val)
                                bound_text = f"{round(ltl, 2)} ~ {round(utl, 2)}"
                                is_pass = (ltl >= limit[0]) and (utl <= limit[1])
                                limit_text = f"{limit[0]} ~ {limit[1]}"
                            else:
                                if operator in ['<=', '<']:
                                    bound_val = mean_val + (k_val_one * std_val)
                                else:
                                    bound_val = mean_val - (k_val_one * std_val)
                                bound_text = round(bound_val, 2)
                                is_pass = check_data_pass(bound_val, operator, limit)
                                limit_text = f"{operator} {limit}"
                            
                            mean_cell = ws.cell(row=mean_row, column=col)
                            mean_cell.value = round(mean_val, 2)
                            is_mean_pass = check_data_pass(mean_val, operator, limit)
                            if not is_mean_pass:
                                mean_cell.fill = pink_fill
                                
                            ws.cell(row=mean_row+1, column=col).value = round(std_val, 2)
                            
                            cell_k = ws.cell(row=mean_row+2, column=col)
                            cell_k.value = bound_text
                            if not is_pass: cell_k.fill = pink_fill
                                
                            cell_p = ws.cell(row=mean_row+3, column=col)
                            cell_p.value = round(p_val, 3)
                            if p_val <= 0.05: cell_p.fill = pink_fill
                                    
                            fig, img_buffer = create_minitab_plot(np.array(data), model_name)
                            img = OpenpyxlImage(img_buffer)
                            img.width = int(9.6 * 37.8)
                            img.height = int(6.2 * 37.8)
                            
                            c_letter = get_column_letter(img_insert_col)
                            ws.add_image(img, f"{c_letter}{header_row}")
                            img_insert_col += 6
                            
                            with cols[col_idx % len(cols)]:
                                st.markdown(f"**{model_name}**")
                                if is_mean_pass:
                                    st.write(f"✅ **평균:** {round(mean_val, 2)}")
                                else:
                                    st.write(f"❌ **평균 불합격:** {round(mean_val, 2)}")
                                
                                if is_pass:
                                    st.success(f"✅ **K값:** {bound_text} (기준: {limit_text})")
                                else:
                                    st.error(f"❌ **K값 불합격:** {bound_text} (기준: {limit_text})")
                                    
                                if p_val > 0.05:
                                    st.info(f"✅ **p-value:** {round(p_val, 3)} (정규성 만족)")
                                else:
                                    st.warning(f"❌ **p-value:** {round(p_val, 3)} (정규성 불만족)")
                                    
                            plt.close(fig) 
                            
                        col_idx += 1
                        
                st.divider() # 대분류 시험(1. Deployment force 등) 하나가 끝날 때마다 구분선 추가

            excel_buffer = io.BytesIO()
            wb.save(excel_buffer)
            excel_buffer.seek(0)

            st.download_button(
                label="📥 결과 엑셀 다운로드 (스마트 판정 적용)",
                data=excel_buffer,
                file_name="스텐트_전체성능분석_완료.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

# =========================================================
# 탭 2: 신규 데이터 생성기 코드 (기존과 동일)
# =========================================================
with tab2:
    st.markdown("### 🎲 원하는 평균/표준편차로 데이터 추출")
    st.info("입력하신 목표치에 100% 일치하도록 보정(Scaling)된 가상의 데이터를 엑셀로 생성합니다.")
    
    col1, col2 = st.columns(2)
    with col1:
        target_mean = st.number_input("🎯 목표 평균", value=10.0, step=0.1)
        target_std = st.number_input("🎯 목표 표준편차", value=0.5, step=0.01)
        decimals = st.number_input("🔢 소수점 자리수", value=2, step=1, min_value=0)
        
    with col2:
        n_cols = st.number_input("가로 (열 개수 / 세트 수)", value=3, step=1, min_value=1)
        n_rows = st.number_input("세로 (행 개수 / 시료 수)", value=10, step=1, min_value=3)

    if st.button("🚀 데이터 추출 및 엑셀 생성", type="primary"):
        data_dict = {}
        for i in range(int(n_cols)):
            raw_data = np.random.normal(loc=target_mean, scale=target_std, size=int(n_rows))
            current_mean = np.mean(raw_data)
            current_std = np.std(raw_data, ddof=1)
            adjusted_data = (raw_data - current_mean) / current_std
            final_data = (adjusted_data * target_std) + target_mean
            final_data = np.round(final_data, int(decimals))
            data_dict[f"Set_{i+1}"] = final_data
            
        df = pd.DataFrame(data_dict)
        
        excel_buffer = io.BytesIO()
        with pd.ExcelWriter(excel_buffer, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Generated_Data')
        excel_buffer.seek(0)
        
        st.success("🎉 데이터 생성이 완료되었습니다! (아래 표에서 미리보기가 가능합니다)")
        st.dataframe(df, use_container_width=True)
        
        st.download_button(
            label="📥 추출된 가상 데이터 엑셀 다운로드",
            data=excel_buffer,
            file_name="가상데이터_추출결과.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
